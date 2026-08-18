# -*- coding: utf-8 -*-
"""검수 엑셀 공용 레이아웃 엔진 — 씬·이미지 통합 A안 (2026-08-14 사용자 확정,
2026-08-18 기미 오일 편 생성기에서 승격).

표준 산출: `씬설계_검수.xlsx` 3시트 (씬·이미지 / 영상 / 공통설정·규칙).
씬·이미지 시트는 왼쪽 씬 정보 열을 그 씬의 이미지 행 수만큼 세로 병합하고
오른쪽에 이미지 정보 열을 행마다 붙인다. 씬~대사 열 틀 고정, 게이트 행
노란 강조, 행 높이는 내용 길이로 계산 (openpyxl은 wrap_text만으로는 행
높이를 안 잡아 첫 줄만 보인다 — 2026-08-14 실측).

편(런)에서는 레이아웃을 복사하지 않는다 — 데이터 파일만 쓴다:

  # output/<run>/ch<N>/씬설계_검수_생성.py (데이터 + 실행 진입점)
  import sys
  from pathlib import Path
  sys.path.insert(0, str(Path(__file__).resolve().parents[3]
                          / ".claude/skills/vox-video/scripts"))
  from make_review_xlsx import FR, build_workbook
  OUT = "output/<run>/ch<N>/씬설계_검수.xlsx"
  SCENES = [...]            # 씬 dict 목록 (아래 계약)
  VIDEO_HEADERS = [...]; VIDEO_ROWS = [...]; VIDEO_WIDTHS = [...]
  COMMON_HEADERS = [...]; COMMON_ROWS = [...]; COMMON_WIDTHS = [...]
  if __name__ == "__main__":
      build_workbook(sys.modules[__name__], OUT)

데이터 계약:
- 씬 dict: id·구간·길이·대사·경계·기법·배속·위험 + frames(FR(...) 목록).
  frames가 비면 병합 안내 1행이 들어간다 — 문구는 씬 dict의
  `빈이미지메모` 키로 지정 (기본 "(자체 이미지 없음)")
- FR(id, 유형, 씬, 역할, 프레이밍, 피사체, 배경, 디테일, 글자, ref,
  선행, 게이트): 게이트가 비어 있지 않으면 그 행이 노랗게 강조된다
- 정본은 scenes/*.md(씬)·manifest.json(합본 트림·램프)이고 엑셀은 뷰다 —
  정본이 바뀐 턴에 데이터 파일을 갱신하고 재실행한다 (scene-design.md,
  SKILL 8단계 동기화 게이트)

자기 검증: `python3 make_review_xlsx.py --selftest`
"""
import math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY = Font(name="Arial", size=10)
GATE_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(*[Side(style="thin", color="D9D9D9")]*4)

# 씬·이미지 시트 표준 열 (통합 A안 — 레이아웃은 엔진이 소유한다)
SCENE_HEADERS = ["씬", "구간(초)", "대사(전체 문장)", "앞 경계", "기법 키",
                 "최대배속", "핵심 위험 · 폴백"]
IMG_HEADERS = ["역할", "ID", "유형", "프레이밍", "피사체 · 상태",
               "배경 · 조명", "디테일", "글자", "ref 첨부 (제외 목록)",
               "선행 조건", "게이트"]
MERGED_WIDTHS = [5, 11, 26, 9, 16, 7, 24] + [10, 5, 6, 36, 36, 26, 30, 11, 24, 14, 13]


def est_lines(text, width):
    """열 너비(엑셀 단위) 안에서 한글 위주 텍스트가 차지할 줄 수 추정.
    한글 1자 ≈ 2단위. 명시적 줄바꿈(\\n)도 센다. 넉넉하게 잡는다 —
    남는 여백이 잘림보다 낫다."""
    if not text:
        return 1
    per_line = max(int(width / 2.1), 4)
    return sum(max(1, math.ceil(len(seg) / per_line)) or 1
               for seg in str(text).split("\n"))


def build_sheet(ws, headers, rows, widths, gate_rows=()):
    ws.freeze_panes = "A2"
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    for r, row in enumerate(rows, 2):
        lines = 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(r, c, v)
            cell.font = BODY; cell.alignment = WRAP; cell.border = THIN
            if r - 2 in gate_rows:
                cell.fill = GATE_FILL
            lines = max(lines, est_lines(v, widths[c - 1]))
        ws.row_dimensions[r].height = min(13.8 * lines + 4, 220)
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def FR(fid, 유형, 씬, 역할, 프레이밍, 피사체="", 배경="", 디테일="", 글자="",
       ref="", 선행="", 게이트=""):
    return dict(id=fid, 유형=유형, 씬=씬, 역할=역할, 프레이밍=프레이밍,
                피사체=피사체, 배경=배경, 디테일=디테일, 글자=글자,
                ref=ref, 선행=선행, 게이트=게이트)


def build_merged(ws, scenes):
    """A안 — 왼쪽 씬 정보 열은 블록(그 씬의 이미지 행들) 높이만큼 세로 병합,
    오른쪽 이미지 정보는 행마다 한 줄. 씬·구간·대사 열은 틀 고정."""
    headers = SCENE_HEADERS + IMG_HEADERS
    widths = MERGED_WIDTHS
    n_sc = len(SCENE_HEADERS)
    n_all = len(headers)
    ws.freeze_panes = "D2"  # 씬·구간·대사 고정 — 오른쪽 스크롤에도 씬 식별 유지
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = HDR_FONT; cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    r = 2
    for s in scenes:
        scene_vals = [s["id"], f"{s['구간']}\n({s['길이']})", s["대사"],
                      s["경계"], s["기법"], s["배속"], s["위험"]]
        # 블록 행들: 프레임 1~2개, 없으면 공유 표시 1행
        frame_rows = [[f["역할"], f["id"], f["유형"], f["프레이밍"], f["피사체"],
                       f["배경"], f["디테일"], f["글자"], f["ref"], f["선행"],
                       f["게이트"]] for f in s["frames"]] or [None]
        gates = [bool(f["게이트"]) for f in s["frames"]] or [False]
        # 행 높이: 이미지 내용 기준으로 잡고, 병합된 씬 열이 더 큰 높이를
        # 요구하면 부족분을 블록 마지막 행에 보탠다
        heights = []
        for row in frame_rows:
            lines = 2 if row is None else max(
                est_lines(v, widths[n_sc + i]) for i, v in enumerate(row))
            heights.append(min(13.8 * lines + 4, 220))
        need = max(13.8 * est_lines(v, widths[i]) + 4
                   for i, v in enumerate(scene_vals))
        if sum(heights) < need:
            heights[-1] += need - sum(heights)
        block = len(frame_rows)
        for k, row in enumerate(frame_rows):
            rr = r + k
            for c in range(1, n_sc + 1):
                cell = ws.cell(rr, c, scene_vals[c - 1] if k == 0 else None)
                cell.font = Font(name="Arial", size=10, bold=(c == 1))
                cell.alignment = WRAP; cell.border = THIN
            if row is None:
                cell = ws.cell(rr, n_sc + 1,
                               s.get("빈이미지메모", "(자체 이미지 없음)"))
                cell.font = BODY; cell.alignment = WRAP
                for c in range(n_sc + 1, n_all + 1):
                    ws.cell(rr, c).border = THIN
                ws.merge_cells(start_row=rr, start_column=n_sc + 1,
                               end_row=rr, end_column=n_all)
            else:
                for i, v in enumerate(row):
                    cell = ws.cell(rr, n_sc + 1 + i, v)
                    cell.font = BODY; cell.alignment = WRAP; cell.border = THIN
                    if gates[k]:
                        cell.fill = GATE_FILL
            ws.row_dimensions[rr].height = heights[k]
        if block > 1:
            for c in range(1, n_sc + 1):
                ws.merge_cells(start_row=r, start_column=c,
                               end_row=r + block - 1, end_column=c)
        r += block
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w


def build_workbook(data, out_path):
    """data: SCENES·VIDEO_*·COMMON_* 속성을 가진 모듈/네임스페이스."""
    wb = Workbook()
    build_merged(wb.active, data.SCENES)
    wb.active.title = "씬·이미지"
    build_sheet(wb.create_sheet("영상"),
                data.VIDEO_HEADERS, data.VIDEO_ROWS, data.VIDEO_WIDTHS)
    build_sheet(wb.create_sheet("공통설정·규칙"),
                data.COMMON_HEADERS, data.COMMON_ROWS, data.COMMON_WIDTHS)
    wb.save(out_path)
    print("저장 완료:", out_path)
    return out_path


# ── 자기 검증 ─────────────────────────────────────────────────
def _selftest():
    import tempfile, os
    from types import SimpleNamespace
    from openpyxl import load_workbook
    data = SimpleNamespace(
        SCENES=[
            dict(id="1", 구간="0~2", 길이="2s", 대사="대사 하나", 경계="시작",
                 기법="기법", 배속="1.0x", 위험="없음",
                 frames=[FR("A", "생성", "1", "first", "프레이밍A"),
                         FR("B", "생성", "1", "end", "프레이밍B",
                            게이트="확인 필요")]),
            dict(id="2", 구간="2~4", 길이="2s", 대사="대사 둘", 경계="체인",
                 기법="공유", 배속="1.0x", 위험="없음", frames=[],
                 빈이미지메모="(씬 1과 파일 공유)"),
        ],
        VIDEO_HEADERS=["클립", "트림"], VIDEO_ROWS=[["clip1", "전체"]],
        VIDEO_WIDTHS=[10, 20],
        COMMON_HEADERS=["항목", "내용"], COMMON_ROWS=[["규칙", "값"]],
        COMMON_WIDTHS=[10, 40],
    )
    fd, path = tempfile.mkstemp(suffix=".xlsx"); os.close(fd)
    try:
        build_workbook(data, path)
        wb = load_workbook(path)
        assert wb.sheetnames == ["씬·이미지", "영상", "공통설정·규칙"], wb.sheetnames
        ws = wb["씬·이미지"]
        assert ws.freeze_panes == "D2"
        merges = {str(m) for m in ws.merged_cells.ranges}
        # 씬 1: 프레임 2행 → 씬 열 7개가 세로 병합
        assert "A2:A3" in merges and "G2:G3" in merges, merges
        # 씬 2: 프레임 없음 → 이미지 열 전체 가로 병합 + 지정 문구
        assert "H4:R4" in merges, merges
        assert ws.cell(4, 8).value == "(씬 1과 파일 공유)"
        # 게이트 행(B, 3행) 노란 강조
        assert ws.cell(3, 8).fill.fgColor.rgb.endswith("FFF2CC")
        assert ws.cell(2, 8).fill.fgColor.rgb != ws.cell(3, 8).fill.fgColor.rgb
        # 행 높이가 계산돼 있다
        assert ws.row_dimensions[2].height and ws.row_dimensions[2].height >= 13.8
        assert wb["영상"].cell(2, 1).value == "clip1"
        assert wb["공통설정·규칙"].freeze_panes == "A2"
        print("selftest PASS — 3시트·세로 병합·빈행 병합·게이트 강조·행 높이")
    finally:
        os.unlink(path)


def _load_data(path):
    import importlib.util, sys as _sys
    from pathlib import Path as _P
    _sys.path.insert(0, str(_P(__file__).resolve().parent))
    spec = importlib.util.spec_from_file_location("review_xlsx_data", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="검수 엑셀 생성 (통합 A안)")
    ap.add_argument("--data", help="SCENES·VIDEO_*·COMMON_* 정의 파이썬 파일")
    ap.add_argument("--out", help="출력 xlsx 경로 (기본: 데이터 파일의 OUT)")
    ap.add_argument("--selftest", action="store_true", help="레이아웃 자기 검증")
    args = ap.parse_args()
    if args.selftest:
        _selftest()
    elif args.data:
        d = _load_data(args.data)
        build_workbook(d, args.out or d.OUT)
    else:
        ap.error("--data 또는 --selftest 필요")
